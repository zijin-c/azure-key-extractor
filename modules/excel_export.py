"""
Excel 导出模块
将 KeyResult 列表导出为 Excel，每个产品一列。
"""

import os
from datetime import datetime
from typing import List

try:
    import openpyxl
    from openpyxl.styles import (
        Font, PatternFill, Alignment, Border, Side
    )
    _OPENPYXL_OK = True
except ImportError:
    _OPENPYXL_OK = False

_RESULTS_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "results")

# 产品列顺序（与 config.PRODUCTS_TO_EXTRACT 一致）
PRODUCT_COLUMNS = [
    "Visio Professional 2021",
    "Project Professional 2021 - DVD",
    "Windows Server 2022 Standard (updated July 2023)",
    "Windows Server 2025 Standard",
    "Windows Server 2019 Standard (updated Mar 2023)",
    "Windows 11 Education, version 25H2",
]

# 列标题简称（Excel 列宽友好）
PRODUCT_SHORT = {
    "Visio Professional 2021":                          "Visio 2021",
    "Project Professional 2021 - DVD":                  "Project 2021",
    "Windows Server 2022 Standard (updated July 2023)": "WinSrv 2022",
    "Windows Server 2025 Standard":                     "WinSrv 2025",
    "Windows Server 2019 Standard (updated Mar 2023)":  "WinSrv 2019",
    "Windows 11 Education, version 25H2":               "Win11 Edu 25H2",
}


def export_to_excel(results, filename: str | None = None) -> str:
    """
    将结果导出为 Excel 文件。
    返回文件路径。
    """
    os.makedirs(_RESULTS_DIR, exist_ok=True)

    if not filename:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"azure_keys_{ts}.xlsx"

    fpath = os.path.join(_RESULTS_DIR, filename)

    if not _OPENPYXL_OK:
        # 降级为 CSV
        return _export_csv(results, fpath.replace(".xlsx", ".csv"))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Azure Keys"

    # ── 样式定义 ──────────────────────────────────────────────
    header_fill   = PatternFill("solid", fgColor="1F3864")
    header_font   = Font(color="FFFFFF", bold=True, size=10)
    success_fill  = PatternFill("solid", fgColor="E2EFDA")
    fail_fill     = PatternFill("solid", fgColor="FCE4D6")
    key_font      = Font(name="Consolas", size=9)
    center_align  = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align    = Alignment(horizontal="left",   vertical="center", wrap_text=True)
    thin_border   = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"),  bottom=Side(style="thin"),
    )

    # ── 表头 ──────────────────────────────────────────────────
    headers = ["#", "时间", "账号 (EDU 邮箱)", "密码", "TOTP Secret", "状态", "备注"]
    for prod in PRODUCT_COLUMNS:
        headers.append(prod)   # 使用产品全称

    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.fill   = header_fill
        cell.font   = header_font
        cell.alignment = center_align
        cell.border = thin_border

    # ── 数据行 ────────────────────────────────────────────────
    for row_idx, result in enumerate(results, 2):
        got = sum(1 for v in result.keys.values() if v)
        row_fill = success_fill if result.success else fail_fill

        base_data = [
            row_idx - 1,
            result.ts,
            result.account.email,
            result.account.password,
            result.totp_secret or "",
            f"✅ {got}/{len(PRODUCT_COLUMNS)}" if result.success else "❌ 失败",
            result.message[:80] if result.message else "",
        ]

        for col_idx, val in enumerate(base_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.fill      = row_fill
            cell.border    = thin_border
            cell.alignment = left_align if col_idx >= 3 else center_align
            if col_idx in (4, 5):  # 密码 / TOTP Secret 用等宽字体
                cell.font = key_font

        # 产品 key 列
        for prod_idx, prod in enumerate(PRODUCT_COLUMNS):
            col_idx = len(base_data) + prod_idx + 1
            key_val = result.keys.get(prod, "")
            cell = ws.cell(row=row_idx, column=col_idx, value=key_val)
            cell.border    = thin_border
            cell.alignment = center_align
            cell.font      = key_font
            if key_val:
                cell.fill = PatternFill("solid", fgColor="D9EAD3")
            else:
                cell.fill = row_fill

    # ── 列宽 ──────────────────────────────────────────────────
    # 产品列加宽（容纳产品全称 + key 长度）
    col_widths = [4, 18, 38, 22, 28, 10, 30] + [42] * len(PRODUCT_COLUMNS)
    for col_idx, width in enumerate(col_widths, 1):
        ws.column_dimensions[
            openpyxl.utils.get_column_letter(col_idx)
        ].width = width

    # 冻结首行
    ws.freeze_panes = "A2"

    # ── 汇总 Sheet ────────────────────────────────────────────
    ws2 = wb.create_sheet("汇总统计")
    ws2.column_dimensions["A"].width = 55
    ws2.column_dimensions["B"].width = 18

    total   = len(results)
    success = sum(1 for r in results if r.success)
    ws2.append(["统计项", "数量"])
    ws2.append(["总账号数", total])
    ws2.append(["成功账号数", success])
    ws2.append(["失败账号数", total - success])
    ws2.append([])
    ws2.append(["产品", "获取到 key 的账号数"])
    for prod in PRODUCT_COLUMNS:
        count = sum(1 for r in results if r.keys.get(prod))
        ws2.append([prod, count])   # 产品全称

    wb.save(fpath)
    return fpath


def _export_csv(results, fpath: str) -> str:
    """降级 CSV 导出（openpyxl 未安装时使用）。"""
    import csv
    headers = ["时间", "账号", "密码", "TOTP Secret", "状态", "备注"] + list(PRODUCT_COLUMNS)
    with open(fpath, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        for r in results:
            row = [
                r.ts, r.account.email, r.account.password, r.totp_secret or "",
                "成功" if r.success else "失败", r.message or "",
            ]
            for prod in PRODUCT_COLUMNS:
                row.append(r.keys.get(prod, ""))
            writer.writerow(row)
    return fpath
